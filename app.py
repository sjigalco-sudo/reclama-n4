import streamlit as st
import pandas as pd
import io
import zipfile
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# Конфигурация страницы
st.set_page_config(page_title="N4 | Full Generator", page_icon="📺", layout="wide")

st.markdown('''
    <style>
    .stButton>button { width: 100%; border-radius: 5px; height: 3em; background-color: #E11D48; color: white; border: none; }
    .stDownloadButton>button { width: 100%; border-radius: 5px; height: 3em; background-color: #2563EB; color: white; border: none; }
    h1 { color: #E11D48; border-bottom: 2px solid #E11D48; padding-bottom: 10px; }
    </style>
    ''', unsafe_allow_html=True)

st.title("📺 N4: ПОЛНЫЙ ГЕНЕРАТОР")

# --- Константы и База Данных ---
DB_FILE = "mp4_database.txt"
PUB_FILE = "PUBLICITATE_HD.mp4"
PUB_DUR = 5.000
SOCIAL_ADS = {
    1: {"file": "1_APA_HD.mpg", "dur": 10.440},
    2: {"file": "2_FRUCTE_HD.mpg", "dur": 10.440},
    3: {"file": "3_MESE HD.mpg", "dur": 10.440},
    4: {"file": "4_MISCARE_HD.mpg", "dur": 10.440},
    5: {"file": "5_SARE_HD.mpg", "dur": 10.440}
}

def load_mp4_ids():
    if os.path.exists(DB_FILE):
        with open(DB_FILE, "r", encoding="utf-8") as f:
            return sorted(list(set([line.strip() for line in f if line.strip()])))
    return ["6856", "7142"]

def save_mp4_ids(id_list):
    cleaned = sorted(list(set([x.strip() for x in id_list if x.strip()])))
    with open(DB_FILE, "w", encoding="utf-8") as f:
        for x in cleaned: f.write(f"{x}\n")
    return cleaned

saved_mp4 = load_mp4_ids()

# --- Вспомогательные функции ---
def xml_escape(text):
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")

def seconds_to_hms_custom(total_seconds):
    total_secs = int(round(total_seconds))
    h = total_secs // 3600
    m = (total_secs % 3600) // 60
    s = total_secs % 60
    return f"{h:02d}:{m:02d}:{s:02d}"

def generate_exact_report(timing_rows, file_date):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.views.sheetView[0].showGridLines = False 
    
    font_regular = Font(name="Calibri", size=11)
    font_bold = Font(name="Calibri", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center", indent=1.0)
    
    blue_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    thin_side = Side(border_style="thin", color="000000")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    ws["E5"] = file_date
    ws["E6"] = "N4"
    ws["E5"].font = font_regular
    ws["E6"].font = font_bold
    
    ws.merge_cells("C7:E7")
    ws["C7"] = "Длительность рекламных блоков"
    ws["C7"].font = font_bold
    ws["C7"].alignment = align_center
    ws["C7"].border = border
    
    ws["C8"] = "Длина ролика"
    ws["D8"] = "" 
    ws["E8"] = "Название блока"
    
    for cell in ["C8", "D8", "E8"]:
        ws[cell].font = font_bold
        ws[cell].alignment = align_center
        ws[cell].border = border
        ws[cell].fill = blue_fill
        
    for idx, row in enumerate(timing_rows, start=9):
        ws.row_dimensions[idx].height = 15.0
        ws[f"C{idx}"] = row["dur"]
        ws[f"C{idx}"].font = font_regular
        ws[f"C{idx}"].alignment = align_center
        ws[f"C{idx}"].border = border
        ws[f"D{idx}"].border = border
        ws[f"E{idx}"] = row["name"]
        ws[f"E{idx}"].font = font_regular
        ws[f"E{idx}"].alignment = align_left
        ws[f"E{idx}"].border = border
        
    ws.column_dimensions['C'].width = 14.5
    ws.column_dimensions['D'].width = 2.5
    ws.column_dimensions['E'].width = 21.0
    
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def to_excel(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    return output.getvalue()

# --- UI ---
with st.sidebar:
    st.header("⚙️ Настройки")
    path_air = st.text_input("Путь к файлам:", value=r"D:\AIR\REKLAMA 2026")
    current_input = st.text_area("ID для MP4:", value=", ".join(saved_mp4))
    if st.button("Сохранить ID"):
        saved_mp4 = save_mp4_ids([x.strip() for x in current_input.split(",") if x.strip()])
        st.toast("💾 База обновлена!")
    uploaded_file = st.file_uploader("Загрузите медиаплан", type=["xls", "xlsx"])

# --- Логика ---
if uploaded_file:
    try:
        base_name = os.path.splitext(uploaded_file.name)[0]
        file_date = next((p for p in base_name.split() if len(p) == 10 and p.count('.') == 2), datetime.now().strftime("%Y-%m-%d"))
        
        df = pd.read_excel(uploaded_file, skiprows=6)
        df_res = df.iloc[:, [2, 7, 9]].copy()
        df_res.columns = ['Block_Time', 'Dur', 'ID']
        df_res['Block_Time'] = pd.to_datetime(df_res['Block_Time'], format='%H:%M:%S', errors='coerce').dt.time.ffill()
        
        grouped = df_res.groupby('Block_Time', sort=False)
        
        zip_buffer = io.BytesIO()
        txt_id_content = io.StringIO()
        xlsx_id_data = []
        timing_rows_formatted = []
        hour_counts = {}
        
        with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
            for i, (block_time, items) in enumerate(grouped, 1):
                h = block_time.hour
                hour_counts[h] = hour_counts.get(h, 0) + 1
                
                report_hour = 24 if h == 0 else h
                report_block_name = f"Реклама {report_hour}.{hour_counts[h]}"
                file_name = f"{report_hour:02d}-{hour_counts[h]}"
                display_time = f"{report_hour}:{block_time.strftime('%M:%S')}"
                
                valid_items = items.dropna(subset=['ID'])
                
                if not valid_items.empty:
                    pure_ads = valid_items['Dur'].sum()
                    soc = SOCIAL_ADS[((i - 1) % 5) + 1]
                    total_dur = PUB_DUR + pure_ads + PUB_DUR + soc['dur']
                    
                    xml = [f'<slblock Source="list" Type="accurate" Sec="{total_dur:.3f}">\nversion 3\n', 
                           f'<item file="{xml_escape(path_air)}\\{PUB_FILE}" dur="{PUB_DUR:.3f}"/>\n']
                    for _, row in valid_items.iterrows():
                        ext = ".mp4" if str(row['ID']).split(".")[0] in saved_mp4 else ".mov"
                        xml.append(f'<item file="{xml_escape(path_air)}\\{str(row["ID"]).split(".")[0]}{ext}" dur="{float(row["Dur"]):.3f}"/>\n')
                    xml.append(f'<item file="{xml_escape(path_air)}\\{PUB_FILE}" dur="{PUB_DUR:.3f}"/>\n')
                    xml.append(f'<item file="{xml_escape(path_air)}\\{soc["file"]}" dur="{soc["dur"]:.3f}"/>\n</slblock>')
                    
                    zip_file.writestr(f"{file_name}.slblock", "".join(xml).encode('utf-16le'))
                    dur_str = seconds_to_hms_custom(total_dur)
                    id_str = "".join([f'"{str(row["ID"]).split(".")[0]}"' for _, row in valid_items.iterrows()])
                else:
                    dur_str, id_str = "00:00:00", ""
                
                timing_rows_formatted.append({"dur": dur_str, "name": report_block_name})
                txt_id_content.write(f"{display_time}\n{id_str}\n\n")
                xlsx_id_data.append({"Время": display_time, "Список ID": id_str})

        st.success("✅ Готово!")
        st.download_button("📥 Скачать блоки", zip_buffer.getvalue(), f"Blocks_{base_name}.zip")
        st.download_button("📥 Скачать отчет (Excel)", generate_exact_report(timing_rows_formatted, file_date), f"Timings_{base_name}.xlsx")
    except Exception as e:
        st.error(f"Ошибка: {e}")
